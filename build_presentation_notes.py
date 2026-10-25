from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os,shutil
R=Path(__file__).resolve().parent
NAVY='183247';TEAL='0E7490';MID='557080'
def setup(title,sub):
 d=Document(); s=d.sections[0]; s.top_margin=Inches(.6);s.bottom_margin=Inches(.6);s.left_margin=Inches(.7);s.right_margin=Inches(.7)
 for name,size,color in [('Normal',9.5,'23313A'),('Title',24,NAVY),('Heading 1',15,NAVY),('Heading 2',11,TEAL),('Heading 3',9,TEAL)]:
  st=d.styles[name];st.font.name='Aptos';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string(color)
  if name!='Normal':st.font.bold=True
 d.sections[0].header.paragraphs[0].text='KRITVA ROBOTICS | KR-MECH-INT-001 | K-EL/01'
 f=d.sections[0].footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=f.add_run('Concept only • 25 October 2026 • Not a manufacturing release');r.font.size=Pt(7);r.font.color.rgb=RGBColor.from_string(MID)
 p=d.add_paragraph(style='Title');p.add_run(title);p=d.add_paragraph();p.add_run(sub).italic=True
 return d
def p(d,t,b=None):
 x=d.add_paragraph();x.paragraph_format.space_after=Pt(4)
 if b and t.startswith(b):x.add_run(b).bold=True;x.add_run(t[len(b):])
 else:x.add_run(t)
def bullet(d,t):
 x=d.add_paragraph(style='List Bullet');x.paragraph_format.space_after=Pt(2);x.add_run(t)
def table(d,h,rows):
 t=d.add_table(rows=1,cols=len(h));t.style='Table Grid'
 for i,v in enumerate(h):t.rows[0].cells[i].text=v
 for c in t.rows[0].cells:
  tcPr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),NAVY);tcPr.append(sh)
  for r in c.paragraphs[0].runs:r.font.color.rgb=RGBColor(255,255,255);r.bold=True;r.font.size=Pt(8)
 rep=OxmlElement('w:tblHeader');rep.set(qn('w:val'),'true');t.rows[0]._tr.get_or_add_trPr().append(rep)
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
 for ri,row in enumerate(rows):
  cells=t.add_row().cells
  t.rows[-1]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for i,v in enumerate(row):
   cells[i].text=v
   for para in cells[i].paragraphs:
    para.paragraph_format.space_after=Pt(0)
    for r in para.runs:r.font.size=Pt(8)
   if ri%2:
    for c in cells:
     sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F2F6F8');c._tc.get_or_add_tcPr().append(sh)
def export(d,name):
 f=R/name;d.save(f);out=R/'lo_export';out.mkdir(exist_ok=True);os.system(f"libreoffice -env:UserInstallation=file:///tmp/lo_note_{os.getpid()} --headless --convert-to pdf --outdir '{out}' '{f}' >/tmp/lo_note.log 2>&1")
 pdf=out/(f.stem+'.pdf')
 if pdf.exists():shutil.copy2(pdf,f.with_suffix('.pdf'))

# Timed notes
d=setup('10-Minute Presentation Script','Use with Kritva_K_EL01_Final_Review_Presentation.pptx; pacing target totals 10:00')
p(d,'Delivery goal: explain the reasoning chain, not simply read numbers. Be candid about what is calculated versus assumed or still awaiting supplier confirmation. Keep the final slide brief enough to leave the review team ready for questions.')
table(d,['Slide','Target time','Speaking script / emphasis'],[
('1 — Title','0:30','“I am presenting K-EL/01, a split-yoke elbow actuator concept for the supplied robotic arm challenge. I focused on making the torque calculation auditable, separating torque transmission from radial support, and calling out what would need confirmation before any hardware release.”'),
('2 — Inputs','1:00','“I took the masses and COM offsets directly from the brief: 1 kg at 350 mm, and 1.5 kg of forearm and hand at 180 mm. The initial configuration is horizontal, so this is a gravity-only static case. Output speed is 45 RPM; BLDC/PMSM, encoder, and planetary-preferred transmission are the stated architecture constraints. The link lengths shown are illustrative, not additional requirements. The brief does not define motion profile, accuracy, impacts or power-loss hold.”'),
('3 — Calculation','1:30','“Using T equals m g r, payload contributes 3.4335 N·m; forearm and hand contribute 2.6487 N·m. Sum equals 6.0822 N·m. I used a preliminary factor of two, giving 12.1644 N·m target. This is an engineering assumption, not a certified safety factor. At 45 RPM, angular velocity is 4.712 rad/s and ideal output power at that target is 57.3 W before losses.”'),
('4 — Drive','1:40','“The candidate is maxon EC-i 52 at 48 V, part 579165, with GP 52 C 91:1 candidate part 223094. Ratio times output speed gives 4095 RPM, only 375 RPM below the motor’s listed nominal speed. That is kinematically plausible but leaves limited headroom. At assumed 40% reducer efficiency, the simplified motor-side screen is 0.334 N·m, 143 W and about 3.64 A using torque constant. At 30%, it fails the simple continuous screen. The motor listing says 180 W class while nominal torque and speed multiply to about 198 W; I will ask the manufacturer for the authoritative rating point and loaded curve. Maxon also notes that EC-i iron-winding motors have a speed-dependent, nonlinear speed/torque characteristic, so I will not use the published gradient as a constant straight-line slope.”'),
('5 — Load path','1:40','“Torque runs from motor, through planetary gearbox and a candidate KTR ROTEX GS flexible coupling, to an independent output spindle, then flange and forearm. Radial loading is carried by the spindle and two spaced 6201 support bearings into the split-yoke and upper arm. That matters because the reducer’s torque rating alone does not establish a cantilever radial-load rating. The CAD bearing centres are 50 mm apart, and the forearm load plane is 50 mm beyond the rear support centre. The STEP model is a simplified envelope assembly, not production geometry.”'),
('6 — Structural / error','1:20','“The two assigned masses give a simplified transverse force of 24.525 N. With a 50 mm overhang, the moment is 1.22625 N·m, explicitly converted to 1226.25 N·mm for the millimetre-based stress equation. A nominal Ø12 mm smooth shaft gives 7.23 MPa bending, 35.85 MPa torsional shear and about 62.52 MPa von Mises. Notch, fatigue, impact and deflection are excluded. One degree of average no-load backlash is about 6.1 mm at the 350 mm payload radius, so precision is an important trade-off.”'),
('7 — Risks / next evidence','1:20','“The first major risk is dynamic and thermal performance. I would obtain the actual trajectory and duty, calculate peak and RMS torque, and get the configured supplier curve and efficiency. The second is output stiffness and load path: verify coupling hubs, spindle geometry, bearing life, fits and carrier stiffness. I would also clarify whether gravity must be held after power loss; there is no brake in this concept, so de-energized testing must be externally supported unless a rated hold solution is later designed.”'),
('8 — Conclusion','1:00','“The concept package supplies the calculation, shortlist, CAD-equivalent STEP and source, labelled concept drawing, BOM, report, traceability, QA and a future validation plan. The assignment does not require buying or building hardware or detailed FEA. FreeCAD was not installed here, so the macro is supplied to generate a native file in FreeCAD. I am not claiming a supplier-confirmed stack, physical test, safety certification or manufacturing release. Thank you; I am ready to discuss the assumptions and closure plan.”')],)
p(d,'Timing note: the eight slide allocations total 10:00. Rehearse with a timer and use the detailed speaking cues rather than reading every word verbatim. Do not accelerate through the calculation or key supplier caveat. If interrupted, preserve the torque math, load path, open gates and concept-only boundary.')
export(d,'Kritva_K_EL01_10_Minute_Presenter_Notes.docx')

# Technical Q&A bank
d=setup('Technical Discussion Q&A Preparation','Suggested structure for the 10-minute follow-on discussion')
p(d,'Answer pattern: state what is known, identify the limit, then name the evidence needed. Do not overstate the concept as validated. A concise engineering answer is usually: calculation → assumption → validation action.')
d.add_heading('Likely questions and model answers',1)
qa=[
('Why did you choose a 2× static torque factor?','It is a preliminary target convention for this assignment, not a formal safety factor. The brief omits acceleration, reflected inertia, impacts, stops, duty and life. I would replace the factor-based target with a trajectory- and risk-derived peak/RMS requirement before selection freeze.'),
('Does the motor actually meet the torque and speed simultaneously?','Not yet proven. The speed ratio implies 4095 RPM, close to the motor’s listed 4470 RPM nominal speed. A loaded torque-speed curve, supply/controller voltage limits, temperature, actual motor/gearhead configuration and drive efficiency are required. The sensitivity arithmetic is screening only.'),
('Why does 180 W differ from nominal torque times speed (~198 W)?','The manufacturer’s product label and the listed nominal torque/speed do not multiply to the same value. I would not speculate about rating convention or curve point. I would ask the vendor to confirm the applicable data sheet, continuous operating point and duty conditions.'),
('Is the published 0.321 rpm/mNm speed/torque gradient a contradictory linear slope?','Not necessarily. Maxon support explains that EC-i iron-winding motors do not have one constant straight-line speed/torque characteristic; the gradient depends on speed. I would request/use the actual motor curve and not extrapolate that single gradient as a constant slope. The separate 180 W label versus nominal Tω arithmetic still needs the applicable rating point clarified.'),
('Why use a gearhead with 1° backlash?','It is a candidate for torque and ratio screening, not an unconditional choice for accurate humanoid motion. At 350 mm, 1° maps to about 6.1 mm arc. If the position-error budget is tighter, use a lower-backlash/stiffer reducer or output feedback, then include compliance and control performance in the full error budget.'),
('Why not let the gearbox output support the forearm?','The gearbox torque rating is not a cantilever radial-load rating. The concept uses an independent spindle supported by two spaced bearings and routes those reactions into the split-yoke and upper arm. Shaft, bearing, fit and carrier calculations are still required.'),
('What changed in your structural screen after the CAD review?','I used the current geometry coordinates: bearing centres at Y=11 and 61 mm and forearm centre plane Y=111 mm. That makes 50 mm spacing and 50 mm overhang. The 1.22625 N·m moment is converted explicitly to 1226.25 N·mm before the mm-based stress formula; the resulting nominal bending stress is 7.23 MPa.'),
('Are your shaft stress numbers a proof of safety?','No. They are nominal static smooth-shaft screens. They exclude keyway/shoulder stress concentrations, cyclic fatigue, impact, manufacturing condition, deflection, actual hub contact and uncertain link loads. A detailed drawing and worst-case dynamic loads are needed.'),
('Why use two 6201 bearings?','They provide a separated support load path in a compact concept. The envelope dimensions align to the nominal 12 mm spindle and 32 mm housing seat. This does not demonstrate fit, oscillatory bearing life, preload, axial loading, lubrication or housing stiffness; those are verification gates.'),
('What are the two biggest requirements you would ask the project lead?','First: full duty/motion profile including acceleration, reversals, impacts, dwell and life because that determines peak/RMS torque and thermal loading. Second: output accuracy/backlash/stiffness and power-loss hold/safe-state requirements because they can change reducer, feedback, brake and controls architecture.'),
('Why no brake?','The brief does not define an unattended gravity-hold requirement. This baseline is a supervised concept and must be externally supported when de-energized. If it must hold the arm after power loss, a normally engaged brake or mechanical counterbalance/support becomes a release requirement and must be sized and safety-reviewed.'),
('What would you test first after a prototype is authorized?','First freeze and verify the configured components and safe-state architecture. Then use a guarded, restrained bench, current limits, E-stop, calibrated torque and RPM measurement, and temperature logging. Start with unloaded and low-load checks, ramp to 45 RPM at defined torque, then run the representative duty cycle and measure backlash/deflection. Predefine abort thresholds.'),
('Did you run FEA or a real prototype test?','No. The brief specifically says prototype and detailed FEA are not required, and no hardware was purchased or made. The package contains a CAD envelope model and a simulation/verification plan; it does not include invented solver output or bench data.'),
('Where is the native FreeCAD file?','FreeCAD was not installed in the generation environment. The deliverable includes a neutral STEP model, CadQuery builder source, and a FreeCAD macro. Open/run the macro inside FreeCAD and save .FCStd locally; the STEP file is the CAD-equivalent exchange deliverable allowed in the brief.'),
('What does the CAD validation tell you?','The STEP re-imported through CadQuery and reported valid geometric solids with 21 solids in the assembly. This checks file geometry validity only. It does not prove no interference, kinematics, tolerance fit, stress, bearing behavior, or manufacturability.'),
('What if the actual efficiency is 30%?','The simplified motor-side requirement rises to about 0.446 N·m and 191 W, which exceeds the listed nominal torque/current screen and 180 W class label. That is a no-go for the simplified continuous requirement unless the actual loaded curve and duty demonstrate acceptable operation; I would investigate another motor/ratio/gearhead and confirm with the supplier.')]
for q,a in qa:
 p(d,q,b=q);p(d,'Answer: '+a)
d.add_heading('Suggested 10-minute discussion pacing',1)
table(d,['Time','Focus','Goal'],[
('0:00–2:00','Torque calculation / margin','Show calculation discipline and assumption boundary.'),('2:00–4:00','Drive selection / motor power caveat','Show manufacturer evidence and no bluffing about open point.'),('4:00–6:00','Load path / shaft support / backlash','Explain why torque and radial support are different.'),('6:00–8:00','Risk, safe state, dynamic/thermal validation','Show awareness of missing system requirements.'),('8:00–10:00','CAD/file status and next steps','State what geometry check proves, and what remains to close.')])
export(d,'Kritva_K_EL01_Technical_Discussion_QA.docx')
